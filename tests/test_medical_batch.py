import os
import io
import time
import pymupdf
import openpyxl
from typing import List

from backend.models.medical_models import MedicalDocumentRecord, PatientInfo
from backend.services.ocr_service import OCRService
from backend.services.medical_extractor import MedicalExtractor
from backend.services.batch_excel_service import BatchExcelService
from backend.services.pdf_service import PDFService

# Helper to generate test PDFs
def create_test_pdf(pages_data: List[dict], rotation: int = 0) -> bytes:
    """
    Creates a synthetic PDF in memory with specified text per page.
    Handles flat [(x, y), text, ...], nested [((x, y), text), ...], and [(x, y, text)].
    """
    doc = pymupdf.open()
    for p_info in pages_data:
        page = doc.new_page(width=595, height=842) # A4
        raw_blocks = p_info.get('text_blocks', [])
        i = 0
        while i < len(raw_blocks):
            item = raw_blocks[i]
            if isinstance(item, tuple) and len(item) == 2 and isinstance(item[0], (int, float)):
                # Flat format: (x, y), "text", (x, y), "text"...
                pt = item
                text = raw_blocks[i + 1] if i + 1 < len(raw_blocks) else ""
                page.insert_text(pt, str(text), fontsize=11)
                i += 2
            elif isinstance(item, (list, tuple)) and len(item) == 2 and isinstance(item[0], (tuple, list)):
                # Nested format: ((x, y), text)
                page.insert_text(item[0], str(item[1]), fontsize=11)
                i += 1
            elif isinstance(item, (list, tuple)) and len(item) == 3:
                # (x, y, text)
                page.insert_text((item[0], item[1]), str(item[2]), fontsize=11)
                i += 1
            else:
                i += 1

        if rotation != 0:
            page.set_rotation(rotation)
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes

def test_01_single_pdf_extraction():
    """Test 1: Single PDF extraction with patient info, assessments, and visit code on 1 page."""
    pdf_bytes = create_test_pdf([{
        'text_blocks': [
            (50, 100), "Patient Name: Robert Smith\nDOB: 05/20/1965\nAge: 59\nGender: Male\nAccount #: 0907-0064\nPatient ID: 0000G230772\nVisit #: V12345",
            (50, 300), "Assessments\n1. Type 2 diabetes mellitus without complications - E11.9\n2. Essential hypertension - I10",
            (50, 450), "Plan:\nContinue metformin 500mg daily.",
            (50, 600), "Visit Code: 99214 Office Visit, Est Pt, Level 4"
        ]
    }])
    ocr = OCRService.get_instance()
    rec = MedicalExtractor.extract_document(pdf_bytes, "test1.pdf", ocr)

    assert rec.success is True
    assert "Robert Smith" in rec.patient_info.patient_name
    assert "05/20/1965" in rec.patient_info.date_of_birth
    assert "0907-0064" in rec.patient_info.account_number
    assert "0000G230772" in rec.patient_info.patient_id
    assert len(rec.assessments) >= 2
    assert "diabetes" in rec.assessments[0].lower()
    assert "99214" in rec.visit_code

def test_03_assessment_on_page_1():
    """Test 3: Assessment located directly on Page 1."""
    pdf_bytes = create_test_pdf([{
        'text_blocks': [
            (50, 100), "Patient Name: Alice Walker\nDOB: 12/04/1980\nAccount #: 001234",
            (50, 250), "Assessments\n1. Acute bronchitis - J20.9",
            (50, 400), "Visit Code: 99203 Office Visit, New Pt, Level 3"
        ]
    }])
    ocr = OCRService.get_instance()
    rec = MedicalExtractor.extract_document(pdf_bytes, "p1.pdf", ocr)
    assert rec.source_pages.get("assessment") == 1
    assert len(rec.assessments) == 1
    assert "bronchitis" in rec.assessments[0].lower()

def test_04_assessment_on_page_2():
    """Test 4: Patient Info on Page 1, Assessment on Page 2."""
    pdf_bytes = create_test_pdf([
        {
            'text_blocks': [(50, 100), "Patient Name: Bob Taylor\nDOB: 03/15/1975\nAccount #: 002345"]
        },
        {
            'text_blocks': [
                (50, 100), "Assessments\n1. Atrial fibrillation - I48.91",
                (50, 300), "Visit Code: 99213 Office Visit"
            ]
        }
    ])
    ocr = OCRService.get_instance()
    rec = MedicalExtractor.extract_document(pdf_bytes, "p2.pdf", ocr)
    assert rec.source_pages.get("patient_information") == 1
    assert rec.source_pages.get("assessment") == 2
    assert "atrial fibrillation" in rec.assessments[0].lower()

def test_05_assessment_on_page_3():
    """Test 5: Patient Info on Page 1, unrelated content on Page 2, Assessment on Page 3."""
    pdf_bytes = create_test_pdf([
        {
            'text_blocks': [(50, 100), "Patient Name: Charlie Brown\nDOB: 08/22/1958\nAccount #: 003456"]
        },
        {
            'text_blocks': [(50, 100), "Past Medical History:\nNo prior surgeries.\nVitals: BP 120/80, Pulse 72"]
        },
        {
            'text_blocks': [
                (50, 100), "Assessments\n1. Chronic kidney disease, stage 3 unspecified - N18.30\n2. History of Lyme disease - Z86.19",
                (50, 300), "Plan:\nNephrology consult."
            ]
        },
        {
            'text_blocks': [(50, 100), "Visit Code: 99204 Office Visit, New Pt., Level 4"]
        }
    ])
    ocr = OCRService.get_instance()
    rec = MedicalExtractor.extract_document(pdf_bytes, "p3.pdf", ocr)
    assert rec.source_pages.get("patient_information") == 1
    assert rec.source_pages.get("assessment") == 3
    assert rec.source_pages.get("visit_code") == 4
    assert len(rec.assessments) == 2
    assert "chronic kidney disease" in rec.assessments[0].lower()

def test_06_assessment_on_page_4():
    """Test 6: Assessment on Page 4 in a 4-page medical report."""
    pdf_bytes = create_test_pdf([
        {'text_blocks': [(50, 100), "Patient Name: Diana Prince\nDOB: 11/11/1985\nAccount #: 004567"]},
        {'text_blocks': [(50, 100), "Social History:\nNon-smoker."]},
        {'text_blocks': [(50, 100), "Review of Systems:\nNormal."]},
        {'text_blocks': [
            (50, 100), "Assessments\n1. Migraine without aura - G43.009",
            (50, 300), "Visit Code: 99215 Comprehensive Visit"
        ]}
    ])
    ocr = OCRService.get_instance()
    rec = MedicalExtractor.extract_document(pdf_bytes, "p4.pdf", ocr)
    assert rec.source_pages.get("patient_information") == 1
    assert rec.source_pages.get("assessment") == 4
    assert "migraine" in rec.assessments[0].lower()

def test_07_assessment_split_across_multiple_lines():
    """Test 7: Assessment item wrapping across multiple lines must be merged into 1 item."""
    pdf_bytes = create_test_pdf([{
        'text_blocks': [
            (50, 100), "Patient Name: Edward Norton\nDOB: 01/01/1970\nAccount #: 005678",
            (50, 200), "Assessments",
            (50, 230), "1. Chronic kidney disease, stage 3",
            (50, 250), "unspecified - N18.30",
            (50, 280), "2. History of Lyme disease - Z86.19",
            (50, 350), "Plan:\nFollow up in 6 months.",
            (50, 450), "Visit Code: 99204 Office Visit"
        ]
    }])
    ocr = OCRService.get_instance()
    rec = MedicalExtractor.extract_document(pdf_bytes, "multiline.pdf", ocr)
    assert len(rec.assessments) == 2
    # Check that line 1 and line 2 were merged into one string
    assert "unspecified - n18.30" in rec.assessments[0].lower()
    assert "chronic kidney disease" in rec.assessments[0].lower()

def test_08_multiple_assessments():
    """Test 8: Extracting multiple numbered assessments (5 items)."""
    pdf_bytes = create_test_pdf([{
        'text_blocks': [
            (50, 100), "Patient Name: Frank Castle\nDOB: 04/14/1972\nAccount #: 006789",
            (50, 200), "Assessments\n1. Hypertension - I10\n2. Hyperlipidemia - E78.5\n3. GERD - K21.9\n4. Type 2 diabetes - E11.9\n5. Vitamin D deficiency - E55.9",
            (50, 400), "Visit Code: 99214 Level 4"
        ]
    }])
    ocr = OCRService.get_instance()
    rec = MedicalExtractor.extract_document(pdf_bytes, "multi_assessments.pdf", ocr)
    assert len(rec.assessments) == 5
    assert "hypertension" in rec.assessments[0].lower()
    assert "vitamin d" in rec.assessments[4].lower()

def test_09_missing_assessment():
    """Test 9: Document without assessment section does not crash and outputs N/A."""
    pdf_bytes = create_test_pdf([{
        'text_blocks': [
            (50, 100), "Patient Name: George Clark\nDOB: 09/09/1990\nAccount #: 007890",
            (50, 300), "Visit Code: 99202 Level 2"
        ]
    }])
    ocr = OCRService.get_instance()
    rec = MedicalExtractor.extract_document(pdf_bytes, "no_assessment.pdf", ocr)
    assert rec.assessments == ["N/A"]
    assert "George Clark" in rec.patient_info.patient_name

def test_10_missing_visit_code():
    """Test 10: Document without visit code sets visit code to N/A."""
    pdf_bytes = create_test_pdf([{
        'text_blocks': [
            (50, 100), "Patient Name: Hannah Abbott\nDOB: 02/02/1988\nAccount #: 008901",
            (50, 250), "Assessments\n1. Seasonal allergic rhinitis - J30.2"
        ]
    }])
    ocr = OCRService.get_instance()
    rec = MedicalExtractor.extract_document(pdf_bytes, "no_visit_code.pdf", ocr)
    assert rec.visit_code == "N/A"
    assert len(rec.assessments) == 1

def test_11_rotated_pdf():
    """Test 11: Document rotated 90 degrees is automatically rotated and extracted."""
    # PyMuPDF set_rotation(90) rotates 90 degrees clockwise
    pdf_bytes = create_test_pdf([{
        'text_blocks': [
            (100, 100), "Patient Name: Ian Malcolm\nDOB: 06/18/1960\nAccount #: 009012",
            (100, 250), "Assessments\n1. Fracture of right femur - S72.91XA",
            (100, 400), "Visit Code: 99205 Complex Emergency"
        ]
    }], rotation=90)
    ocr = OCRService.get_instance()
    rec = MedicalExtractor.extract_document(pdf_bytes, "rotated.pdf", ocr)
    assert "Ian Malcolm" in rec.patient_info.patient_name or "06/18/1960" in rec.patient_info.date_of_birth
    assert len(rec.assessments) >= 1

def test_12_ocr_errors_and_fuzzy_headings():
    """Test 12: OCR imperfections like 'ASSESSMENT:' and 'Visit C0de' are correctly recognized."""
    pdf_bytes = create_test_pdf([{
        'text_blocks': [
            (50, 100), "Patient: Julia Roberts\nBirth Date: 10/28/1967\nAcct #: 010123",
            (50, 250), "ASSESSMENT:\n1) Chronic sinusitis - J32.9",
            (50, 400), "Visit C0de:\n99213 Office Visit"
        ]
    }])
    ocr = OCRService.get_instance()
    rec = MedicalExtractor.extract_document(pdf_bytes, "ocr_errors.pdf", ocr)
    assert "Julia Roberts" in rec.patient_info.patient_name
    assert len(rec.assessments) >= 1
    assert "sinusitis" in rec.assessments[0].lower()
    assert "99213" in rec.visit_code

def test_13_leading_zeros_preservation(tmp_path):
    """Test 13: Excel generation strictly preserves leading zeros (0907-0064, 0000G230772)."""
    rec = MedicalDocumentRecord(
        source_file="patient_zero.pdf",
        patient_info=PatientInfo(
            patient_name="Zero Test",
            date_of_birth="07/15/1951",
            account_number="0907-0064",
            patient_id="0000G230772",
            visit_number="0001928"
        ),
        assessments=["1. Test condition - T01.0"],
        visit_code="0099204 Office Visit"
    )
    headers, rows, low_conf = BatchExcelService.build_headers_and_rows([rec])
    out_xlsx = str(tmp_path / "test_zeros.xlsx")
    BatchExcelService.create_batch_excel_file(headers, rows, out_xlsx)

    wb = openpyxl.load_workbook(out_xlsx)
    ws = wb["Table 1"]
    
    # Locate column indices
    header_vals = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    acct_col = header_vals.index("Account Number") + 1
    pat_id_col = header_vals.index("Patient ID") + 1
    visit_num_col = header_vals.index("Visit Number") + 1

    # Check cell values and string types in row 2
    cell_acct = ws.cell(2, acct_col)
    cell_pid = ws.cell(2, pat_id_col)
    cell_vnum = ws.cell(2, visit_num_col)

    assert str(cell_acct.value) == "0907-0064"
    assert str(cell_pid.value) == "0000G230772"
    assert str(cell_vnum.value) == "0001928"
    assert cell_acct.data_type == 's'
    assert cell_pid.data_type == 's'

def test_14_one_failed_pdf_inside_batch(tmp_path):
    """Test 14: One failed/corrupted PDF in a batch does NOT crash the batch."""
    good_pdf = create_test_pdf([{
        'text_blocks': [
            (50, 100), "Patient Name: Good Patient\nDOB: 01/01/2000\nAccount #: 001122",
            (50, 250), "Assessments\n1. Healthy examination - Z00.00",
            (50, 400), "Visit Code: 99381 Preventive Care"
        ]
    }])
    bad_pdf = b"NOT A VALID PDF CONTENT"

    ocr = OCRService.get_instance()
    rec1 = MedicalExtractor.extract_document(good_pdf, "good.pdf", ocr)
    rec2 = MedicalExtractor.extract_document(bad_pdf, "corrupted.pdf", ocr)

    assert rec1.success is True
    assert rec2.success is False
    assert rec2.error is not None

    # Excel should still generate successfully with both records
    headers, rows, low_conf = BatchExcelService.build_headers_and_rows([rec1, rec2])
    assert len(rows) == 2
    out_xlsx = str(tmp_path / "batch_with_error.xlsx")
    BatchExcelService.create_batch_excel_file(headers, rows, out_xlsx)
    assert os.path.exists(out_xlsx)

def test_02_and_15_ten_pdf_batch_producing_ten_excel_rows(tmp_path):
    """
    Test 2 & Test 15:
    10 PDFs batch processed sequentially producing exactly ONE Excel file
    with exactly 10 rows (1 row per PDF) and dynamic Assessment columns.
    """
    records = []
    ocr = OCRService.get_instance()

    for i in range(1, 11):
        # Varying number of assessments across documents (1 to 4)
        num_assessments = (i % 4) + 1
        assessments_text = "\n".join([f"{a}. Diagnosis {a} for patient {i} - D0{a}.0" for a in range(1, num_assessments + 1)])

        pdf_bytes = create_test_pdf([{
            'text_blocks': [
                (50, 100), f"Patient Name: Patient {i:02d}\nDOB: 0{i}/10/1980\nAccount #: 0090-{i:04d}\nPatient ID: PID{i:05d}",
                (50, 250), f"Assessments\n{assessments_text}",
                (50, 500), f"Visit Code: 9920{i%5} Office Visit"
            ]
        }])
        rec = MedicalExtractor.extract_document(pdf_bytes, f"patient_{i:02d}.pdf", ocr)
        records.append(rec)

    assert len(records) == 10
    for r in records:
        assert r.success is True

    # Build Excel
    headers, table_data, low_conf = BatchExcelService.build_headers_and_rows(records)

    # Max assessments is 4, so headers should have Assessment 1 through Assessment 4
    assert "Assessment 1" in headers
    assert "Assessment 4" in headers
    assert len(table_data) == 10  # Exactly 10 rows!

    out_xlsx = str(tmp_path / "ten_patients_batch.xlsx")
    BatchExcelService.create_batch_excel_file(headers, table_data, out_xlsx)

    # Verify Excel workbook properties
    wb = openpyxl.load_workbook(out_xlsx)
    assert "Table 1" in wb.sheetnames
    ws = wb["Table 1"]
    
    # 1 header row + 10 data rows = 11 rows total
    assert ws.max_row == 11
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None

    # Check that row 10 has Source PDF
    source_pdf_val = ws.cell(11, 1).value
    assert source_pdf_val == "patient_10.pdf"
